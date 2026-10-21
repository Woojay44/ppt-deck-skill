"""Structural QA for an exported academic deck.

Usage: python qa_pptx.py DECK.pptx [--min-pt 18] [--allow-small "数据来源,资料来源,/ 15"]

Checks every slide for: text runs below the minimum size (except allowed
footer/page-number patterns), fonts actually used (theme fonts resolved),
speaker notes present, picture count, and empty slides. Exit code 1 when a
blocking issue is found.
"""
import argparse
import re
import sys
from collections import Counter

from pptx import Presentation
from pptx.util import Pt


def theme_fonts(prs):
    """Resolve +mn-/+mj- theme font references from the first slide master."""
    xml = prs.slide_masters[0].part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    ).blob.decode("utf-8", "ignore")
    out = {}
    for kind, tag in (("mn", "minorFont"), ("mj", "majorFont")):
        m = re.search(rf"<a:{tag}>(.*?)</a:{tag}>", xml, re.S)
        if m:
            for script in ("latin", "ea"):
                f = re.search(rf'<a:{script} typeface="([^"]*)"', m.group(1))
                if f:
                    out[f"+{kind}-{'lt' if script == 'latin' else 'ea'}"] = f.group(1)
    return out


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--min-pt", type=float, default=18)
    ap.add_argument("--allow-small", default="数据来源,资料来源, / ")
    a = ap.parse_args()
    allow = [s for s in a.allow_small.split(",") if s]

    prs = Presentation(a.pptx)
    tf = theme_fonts(prs)
    fonts, blocking, notes_missing, small_allowed = Counter(), [], [], []
    print(f"deck: {a.pptx}  slides: {len(prs.slides)}  theme fonts: {tf}")
    for i, slide in enumerate(prs.slides, 1):
        n_text = n_pic = 0
        for sh in iter_shapes(slide.shapes):
            if sh.shape_type == 13:
                n_pic += 1
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if not text:
                    continue
                n_text += 1
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    rpr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    for tag in ("latin", "ea"):
                        el = rpr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}") if rpr is not None else None
                        face = el.get("typeface") if el is not None else f"+mn-{'lt' if tag == 'latin' else 'ea'}"
                        fonts[tf.get(face, face)] += 1
                    size = r.font.size or (para.font.size if para.font.size else None)
                    if size is not None and size < Pt(a.min_pt):
                        rec = f"slide {i}: {size.pt:.1f}pt «{text[:30]}»"
                        is_page_no = re.fullmatch(r"\d+(\s*/\s*\d+)?", text) is not None
                        (small_allowed if is_page_no or any(p in text for p in allow) else blocking).append(rec)
        has_notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        if not has_notes:
            notes_missing.append(i)
        print(f"  slide {i:2d}: text paragraphs {n_text:3d}  pictures {n_pic}  notes {'yes' if has_notes else 'NO'}")
        if n_text == 0 and n_pic == 0:
            blocking.append(f"slide {i}: empty slide")
    print("fonts used:", dict(fonts))
    print(f"allowed small text (footer/page number): {len(small_allowed)}")
    for r in blocking:
        print("  [BLOCK]", r)
    if notes_missing:
        print("  [WARN] slides without notes:", notes_missing)
    print("RESULT:", "FAIL" if blocking else "PASS")
    sys.exit(1 if blocking else 0)


if __name__ == "__main__":
    main()
